import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # Color Palette
    BG_DARK = RGBColor(11, 15, 25)        # #0b0f19
    CARD_BG = RGBColor(19, 26, 43)        # #131a2b
    CARD_BORDER = RGBColor(38, 51, 82)    # #263352
    BRAND_BLUE = RGBColor(59, 130, 246)   # #3b82f6
    CYAN_ACCENT = RGBColor(6, 182, 212)   # #06b6d4
    EMERALD = RGBColor(16, 185, 129)      # #10b981
    AMBER = RGBColor(245, 158, 11)        # #f59e0b
    PURPLE = RGBColor(168, 85, 247)       # #a855f7
    TEXT_WHITE = RGBColor(255, 255, 255)
    TEXT_MUTED = RGBColor(148, 163, 184)  # #94a3b8

    blank_layout = prs.slide_layouts[6]

    def set_slide_background(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
        bg.fill.solid()
        bg.fill.fore_color.rgb = BG_DARK
        bg.line.fill.background()

    def add_header(slide, badge_text, title_text, subtitle_text=""):
        # Badge Pill
        badge = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.5), Inches(3.2), Inches(0.35))
        badge.fill.solid()
        badge.fill.fore_color.rgb = RGBColor(30, 58, 138) # brand blue tint
        badge.line.color.rgb = BRAND_BLUE
        tf = badge.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = badge_text.upper()
        p.font.size = Pt(10)
        p.font.bold = True
        p.font.color.rgb = CYAN_ACCENT
        p.alignment = PP_ALIGN.CENTER

        # Main Title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.9), Inches(11.7), Inches(0.8))
        tf = title_box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title_text
        p.font.size = Pt(26)
        p.font.bold = True
        p.font.color.rgb = TEXT_WHITE

        # Subtitle
        if subtitle_text:
            sub_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.65), Inches(11.7), Inches(0.45))
            tf = sub_box.text_frame
            p = tf.paragraphs[0]
            p.text = subtitle_text
            p.font.size = Pt(13)
            p.font.color.rgb = TEXT_MUTED

    def add_card(slide, left, top, width, height, title, items, badge="", accent_color=BRAND_BLUE):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = CARD_BORDER
        card.line.width = Pt(1.5)

        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.25)
        tf.margin_right = Inches(0.25)
        tf.margin_top = Inches(0.25)
        tf.margin_bottom = Inches(0.2)

        # Header in card
        p0 = tf.paragraphs[0]
        if badge:
            p0.text = f"[{badge}] {title}"
        else:
            p0.text = title
        p0.font.size = Pt(16)
        p0.font.bold = True
        p0.font.color.rgb = accent_color
        p0.space_after = Pt(10)

        for item in items:
            p = tf.add_paragraph()
            p.text = f"•  {item}"
            p.font.size = Pt(12)
            p.font.color.rgb = TEXT_MUTED
            p.space_after = Pt(6)

    # -------------------------------------------------------------------------
    # SLIDE 1: Title & Hero
    # -------------------------------------------------------------------------
    s1 = prs.slides.add_slide(blank_layout)
    set_slide_background(s1)

    # Big decorative glow card
    glow = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.0), Inches(11.733), Inches(5.5))
    glow.fill.solid()
    glow.fill.fore_color.rgb = CARD_BG
    glow.line.color.rgb = BRAND_BLUE
    glow.line.width = Pt(2)

    gtf = glow.text_frame
    gtf.word_wrap = True
    gtf.margin_left = Inches(0.6)
    gtf.margin_top = Inches(0.6)

    p = gtf.paragraphs[0]
    p.text = "CAMPUSGET ENTERPRISE PRESENTATION"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = CYAN_ACCENT
    p.space_after = Pt(12)

    p = gtf.add_paragraph()
    p.text = "Enterprise School ERP &\nModern Public Website"
    p.font.size = Pt(38)
    p.font.bold = True
    p.font.color.rgb = TEXT_WHITE
    p.space_after = Pt(16)

    p = gtf.add_paragraph()
    p.text = "A Complete Multi-Tenant Platform for Institutional Operations, Admissions, Fee Accounting, and Student Engagement."
    p.font.size = Pt(16)
    p.font.color.rgb = TEXT_MUTED
    p.space_after = Pt(24)

    p = gtf.add_paragraph()
    p.text = "🚀 Live Demo: https://stxaviers-school-erp-bycampusget.netlify.app\n💻 Tech Stack: React 18 • TypeScript • Node.js • Express • PostgreSQL 16 • Redis • Docker\n👤 Author: Sandeep Gaud | CampusGet Verified Package"
    p.font.size = Pt(13)
    p.font.color.rgb = RGBColor(165, 180, 252) # soft indigo

    # -------------------------------------------------------------------------
    # SLIDE 2: Problem & Solution
    # -------------------------------------------------------------------------
    s2 = prs.slides.add_slide(blank_layout)
    set_slide_background(s2)
    add_header(s2, "Executive Summary", "The Challenge vs. The Enterprise Solution", "Transforming fragmented manual processes into an integrated, paperless school management ecosystem.")

    add_card(s2, Inches(0.8), Inches(2.2), Inches(5.6), Inches(4.7), 
             "Traditional School Challenges", [
                 "Paper attendance registers prone to loss and manual calculation errors.",
                 "Un-reconciled fee collections, delayed receipt generation, and payment leakages.",
                 "Disconnected communications across notice boards, SMS, and scattered emails.",
                 "Manual exam marks tabulation leading to delayed report card generation.",
                 "Lacks transparent parent engagement and real-time ward monitoring."
             ], badge="PAIN POINTS", accent_color=RGBColor(239, 68, 68))

    add_card(s2, Inches(6.8), Inches(2.2), Inches(5.7), Inches(4.7), 
             "The Enterprise ERP Platform", [
                 "Unified Cloud Monorepo connecting Administration, Teachers, Students & Parents.",
                 "Automated Double-Entry Fee Ledger with quarterly/annual transparent calculators.",
                 "1-Click Section Attendance with instant daily summaries and leave tracking.",
                 "Automated Examination Engine with grading scales and weighted report cards.",
                 "Modern SEO-Ready Public Website with online admissions enquiry capture."
             ], badge="THE SOLUTION", accent_color=EMERALD)

    # -------------------------------------------------------------------------
    # SLIDE 3: Public Website & Modern UI
    # -------------------------------------------------------------------------
    s3 = prs.slides.add_slide(blank_layout)
    set_slide_background(s3)
    add_header(s3, "Frontend Experience", "High-Conversion Public Website & Modern UX", "Designed to deliver a world-class institutional image and maximize student admissions.")

    add_card(s3, Inches(0.8), Inches(2.2), Inches(3.7), Inches(4.7), 
             "Institutional Branding", [
                 "Modern Hero Section with animated typography and dynamic call-to-actions.",
                 "Principal's Welcome Desk showcasing institutional vision and accreditation.",
                 "Academics Program Wings from Primary to Senior Secondary Specializations.",
                 "Interactive Campus Infrastructure photo gallery and event highlights."
             ], badge="WEBSITE CMS", accent_color=CYAN_ACCENT)

    add_card(s3, Inches(4.8), Inches(2.2), Inches(3.7), Inches(4.7), 
             "Day & Night Mode Engine", [
                 "Fully functional Light (Day) and Dark (Night) mode switch.",
                 "Smooth 180° animated icon toggle with sun/moon transitions.",
                 "Persistent user preference preserved across browser sessions via localStorage.",
                 "High-contrast custom educational color palette tailored for accessibility."
             ], badge="THEME ENGINE", accent_color=AMBER)

    add_card(s3, Inches(8.8), Inches(2.2), Inches(3.7), Inches(4.7), 
             "Lead & Admissions Capture", [
                 "Direct Admissions Enquiry form with real-time field validation.",
                 "Visitor contact routing and campus tour booking mechanisms.",
                 "SEO-optimized semantic markup with high Lighthouse performance score.",
                 "Fully responsive layout optimized for smartphones, tablets, and laptops."
             ], badge="ENGAGEMENT", accent_color=EMERALD)

    # -------------------------------------------------------------------------
    # SLIDE 4: Transparent Fee Structure
    # -------------------------------------------------------------------------
    s4 = prs.slides.add_slide(blank_layout)
    set_slide_background(s4)
    add_header(s4, "Financial Transparency", "Public Fee Structure & Financial Engine", "Eliminating parent hesitation with 100% upfront pricing and institutional aid policies.")

    add_card(s4, Inches(0.8), Inches(2.2), Inches(5.6), Inches(4.7), 
             "Public Fee Calculator", [
                 "Annual vs. Quarterly Payment installment toggle.",
                 "Grade Wing Breakdown: Primary, Middle, Secondary, Senior Secondary.",
                 "Transparent itemization of tuition fees, activity kits, and admission charges.",
                 "Inclusive amenities checklist (Robotics labs, digital vault, sports coaching).",
                 "Institutional Scholarships: Merit waivers (50%), Sibling concessions (15%), and Sports quotas (40%)."
             ], badge="PARENT TRANSPARENCY", accent_color=CYAN_ACCENT)

    add_card(s4, Inches(6.8), Inches(2.2), Inches(5.7), Inches(4.7), 
             "ERP Finance Engine", [
                 "Double-Entry Accounting principles ensuring ledger integrity.",
                 "Bulk Invoicing Engine generating term-wise billings automatically.",
                 "Idempotent Payment Collection protected via Idempotency-Key headers.",
                 "Instant digital receipt generation and transaction audit trails.",
                 "Complete payment mode support: UPI, Netbanking, Cards, and Zero-Cost EMI."
             ], badge="ACCOUNTING ENGINE", accent_color=PURPLE)

    # -------------------------------------------------------------------------
    # SLIDE 5: Academic Operations
    # -------------------------------------------------------------------------
    s5 = prs.slides.add_slide(blank_layout)
    set_slide_background(s5)
    add_header(s5, "Core Operations", "Academic Operations & Examination Engine", "Streamlining daily classroom workflows, attendance recording, and grade evaluations.")

    add_card(s5, Inches(0.8), Inches(2.2), Inches(3.7), Inches(4.7), 
             "Attendance Management", [
                 "Section-wise bulk attendance recording (Present, Absent, Late, Half-Day).",
                 "Automated student attendance percentage calculations.",
                 "Real-time notifications sent to guardians for absent wards.",
                 "Formal leave application submission and teacher approval workflow."
             ], badge="ATTENDANCE", accent_color=EMERALD)

    add_card(s5, Inches(4.8), Inches(2.2), Inches(3.7), Inches(4.7), 
             "Exams & Report Cards", [
                 "Term exam scheduling with subject and classroom allocations.",
                 "Grid marks entry interface allowing rapid bulk score input.",
                 "Automated grade scale mapping (A+, A, B, C) and weighted GPA calculation.",
                 "One-click compiled printable student report cards with teacher remarks."
             ], badge="EVALUATIONS", accent_color=BRAND_BLUE)

    add_card(s5, Inches(8.8), Inches(2.2), Inches(3.7), Inches(4.7), 
             "Timetable & Homework", [
                 "Slot allocation engine with built-in conflict detection preventing double-booking.",
                 "Teacher weekly workload summary and room utilization metrics.",
                 "Digital homework assignment creation with due-date enforcement.",
                 "Student submission upload handler with teacher grading grid."
             ], badge="CLASSROOM", accent_color=PURPLE)

    # -------------------------------------------------------------------------
    # SLIDE 6: Multi-Role Portals & RBAC
    # -------------------------------------------------------------------------
    s6 = prs.slides.add_slide(blank_layout)
    set_slide_background(s6)
    add_header(s6, "Security & Governance", "Role-Based Access Control (RBAC) & Portals", "Personalized dashboards and strict permission enforcement for all 8 institutional roles.")

    add_card(s6, Inches(0.8), Inches(2.2), Inches(3.7), Inches(2.3), 
             "School Administrator", [
                 "Master system configuration & multi-school scope.",
                 "Academic year setup, staff directory & admission ledger.",
                 "Audit log inspection and institutional analytics."
             ], badge="SUPER ADMIN", accent_color=BRAND_BLUE)

    add_card(s6, Inches(4.8), Inches(2.2), Inches(3.7), Inches(2.3), 
             "Teacher Portal", [
                 "Assigned classes overview & weekly schedule grid.",
                 "Rapid student attendance entry and roster view.",
                 "Homework assignment dispatcher & exam marks input."
             ], badge="FACULTY", accent_color=EMERALD)

    add_card(s6, Inches(8.8), Inches(2.2), Inches(3.7), Inches(2.3), 
             "Parent Console", [
                 "Multi-child ward switcher for seamless monitoring.",
                 "Real-time attendance rates, leave requests & fee dues.",
                 "Direct exam report card downloads & school notices."
             ], badge="GUARDIANS", accent_color=AMBER)

    add_card(s6, Inches(0.8), Inches(4.7), Inches(3.7), Inches(2.3), 
             "Student Self-Service", [
                 "Personal weekly class schedule & upcoming exams.",
                 "Homework assignment downloads and submissions.",
                 "Historical report cards and activity notice board."
             ], badge="STUDENTS", accent_color=CYAN_ACCENT)

    add_card(s6, Inches(4.8), Inches(4.7), Inches(3.7), Inches(2.3), 
             "Accountant & Finance", [
                 "Fee structure configuration & student invoices.",
                 "Point-of-sale fee collection & instant receipts.",
                 "Audit refund records & daily revenue reconciliations."
             ], badge="FINANCE", accent_color=PURPLE)

    add_card(s6, Inches(8.8), Inches(4.7), Inches(3.7), Inches(2.3), 
             "Instant Demo Fallback", [
                 "Built-in offline client simulation for testing.",
                 "1-Click role preview without database setup.",
                 "Seamlessly activates live PostgreSQL when online."
             ], badge="DEMO ENGINE", accent_color=RGBColor(244, 63, 94))

    # -------------------------------------------------------------------------
    # SLIDE 7: Advanced Modules
    # -------------------------------------------------------------------------
    s7 = prs.slides.add_slide(blank_layout)
    set_slide_background(s7)
    add_header(s7, "Comprehensive Suite", "Advanced Specialized School Modules", "Beyond core academics: Complete facility and administrative lifecycle management.")

    add_card(s7, Inches(0.8), Inches(2.2), Inches(2.75), Inches(4.7), 
             "Library System", [
                 "Catalog search by title, author, and ISBN.",
                 "Transactional book issue and return tracking.",
                 "Real-time available copies monitor.",
                 "Automated overdue fine computation.",
                 "Student borrow history records."
             ], badge="LIBRARY", accent_color=CYAN_ACCENT)

    add_card(s7, Inches(3.75), Inches(2.2), Inches(2.75), Inches(4.7), 
             "Transport Fleet", [
                 "Vehicle fleet registration and compliance.",
                 "Driver assignment & route fare structures.",
                 "Bus stop allocation with pickup timings.",
                 "Vehicle seating capacity enforcement.",
                 "Student commute safety rosters."
             ], badge="LOGISTICS", accent_color=AMBER)

    add_card(s7, Inches(6.7), Inches(2.2), Inches(2.75), Inches(4.7), 
             "HR & Payroll", [
                 "Employee directory and contracts.",
                 "Salary components: Basic, HRA, Allowances.",
                 "Statutory tax and provident deductions.",
                 "1-Click monthly bulk payroll runs.",
                 "Downloadable digital payslips."
             ], badge="PAYROLL", accent_color=EMERALD)

    add_card(s7, Inches(9.65), Inches(2.2), Inches(2.8), Inches(4.7), 
             "Certificates & Stock", [
                 "Dynamic Certificate Generator (TC, Bonafide).",
                 "Token interpolation: {{name}}, {{admissionNo}}.",
                 "Inventory tracking with SKU movements.",
                 "Supplier directory and purchase orders.",
                 "Immutable security audit logs."
             ], badge="CERTIFICATES", accent_color=PURPLE)

    # -------------------------------------------------------------------------
    # SLIDE 8: System Architecture & Tech Stack
    # -------------------------------------------------------------------------
    s8 = prs.slides.add_slide(blank_layout)
    set_slide_background(s8)
    add_header(s8, "Engineering Excellence", "System Architecture & Production Tech Stack", "Built with modern, scalable, enterprise-grade technologies for maximum reliability.")

    add_card(s8, Inches(0.8), Inches(2.2), Inches(5.6), Inches(4.7), 
             "Architecture & Security", [
                 "Clean Monorepo: apps/web, apps/api, packages/shared.",
                 "Discriminator-Column Multi-Tenancy (organizationId, schoolId).",
                 "Stateless JWT Access Tokens (15m) + HttpOnly Refresh Tokens (7d).",
                 "Redis Session Whitelist supporting instant session revocation.",
                 "Database Indexing optimized for sub-millisecond query execution.",
                 "Strict Zod schema validation on all API boundaries.",
                 "Audit Logging: captures user, IP, action, and state changes."
             ], badge="ENTERPRISE CORE", accent_color=BRAND_BLUE)

    add_card(s8, Inches(6.8), Inches(2.2), Inches(5.7), Inches(4.7), 
             "Production Technology Stack", [
                 "Frontend: React 18, Vite, TypeScript, Tailwind CSS, Lucide.",
                 "Backend: Node.js, Express, TypeScript, Prisma ORM.",
                 "Database: PostgreSQL 16 with relational integrity.",
                 "Cache: Redis 7 for high-performance session validation.",
                 "Containerization: Docker & docker-compose.prod.yml.",
                 "Frontend CDN: Netlify-ready with 1-click CI/CD deployment.",
                 "Performance: k6 tested for 200+ concurrent virtual users."
             ], badge="TECH SPECIFICATIONS", accent_color=EMERALD)

    # -------------------------------------------------------------------------
    # SLIDE 9: Business Value & ROI
    # -------------------------------------------------------------------------
    s9 = prs.slides.add_slide(blank_layout)
    set_slide_background(s9)
    add_header(s9, "Commercial Value", "Institutional Impact & Return on Investment", "Tangible operational savings and heightened reputation for educational institutions.")

    add_card(s9, Inches(0.8), Inches(2.2), Inches(3.7), Inches(4.7), 
             "85% Paper Reduction", [
                 "Replaces physical attendance registers with instant digital recording.",
                 "Automated electronic report cards eliminate printing costs.",
                 "Online admission enquiry pipeline replaces physical paperwork.",
                 "Digital receipts and fee ledgers reduce billing stationery."
             ], badge="EFFICIENCY", accent_color=EMERALD)

    add_card(s9, Inches(4.8), Inches(2.2), Inches(3.7), Inches(4.7), 
             "Zero Fee Leakage", [
                 "Double-entry ledger rules ensure every rupee is accounted for.",
                 "Real-time outstanding dues alert system for accountant follow-ups.",
                 "Transparent fee structure builds parent trust and timely collections.",
                 "Audit trails prevent unauthorized fee modifications."
             ], badge="REVENUE", accent_color=CYAN_ACCENT)

    add_card(s9, Inches(8.8), Inches(2.2), Inches(3.7), Inches(4.7), 
             "Parent Satisfaction", [
                 "Parents access real-time attendance, fee receipts, and report cards.",
                 "Instant school announcements and emergency event notifications.",
                 "Seamless multi-child management under one unified login.",
                 "Enhances school prestige with a modern digital ecosystem."
             ], badge="SATISFACTION", accent_color=PURPLE)

    # -------------------------------------------------------------------------
    # SLIDE 10: Conclusion & Live Demonstration
    # -------------------------------------------------------------------------
    s10 = prs.slides.add_slide(blank_layout)
    set_slide_background(s10)

    cta_card = s10.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.0), Inches(11.733), Inches(5.5))
    cta_card.fill.solid()
    cta_card.fill.fore_color.rgb = CARD_BG
    cta_card.line.color.rgb = CYAN_ACCENT
    cta_card.line.width = Pt(2)

    ctf = cta_card.text_frame
    ctf.word_wrap = True
    ctf.margin_left = Inches(0.8)
    ctf.margin_top = Inches(0.8)

    p = ctf.paragraphs[0]
    p.text = "VERIFIED PRODUCTION PACKAGE"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = CYAN_ACCENT
    p.space_after = Pt(14)

    p = ctf.add_paragraph()
    p.text = "Ready for Institutional Deployment\n& Client Demonstrations"
    p.font.size = Pt(36)
    p.font.bold = True
    p.font.color.rgb = TEXT_WHITE
    p.space_after = Pt(20)

    p = ctf.add_paragraph()
    p.text = "🔗 Live Working Application: https://stxaviers-school-erp-bycampusget.netlify.app"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = BRAND_BLUE
    p.space_after = Pt(10)

    p = ctf.add_paragraph()
    p.text = "📁 Verified GitHub Source: https://github.com/Sandeep798512/school-erp"
    p.font.size = Pt(15)
    p.font.color.rgb = TEXT_MUTED
    p.space_after = Pt(24)

    p = ctf.add_paragraph()
    p.text = "Thank You! • Questions & Discussion Welcome • Powered by CampusGet"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = EMERALD

    # Save presentation
    output_path = r"C:\Users\Er. SANDEEP GAUD\Desktop\School_ERP_Presentation.pptx"
    prs.save(output_path)
    print(f"Presentation successfully generated and saved to: {output_path}")

if __name__ == "__main__":
    create_presentation()
